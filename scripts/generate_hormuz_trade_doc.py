from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


OUTPUT_PATH = Path(
    "/Users/timli/Code/autoresearch-equity/deliverables/hormuz_fresh_screen_trade_ideas_2026-03-26.docx"
)


MACRO_POINTS = [
    "The usable macro signal is real, but the duration is uncertain. The IEA says more than 110 bcm of LNG, about one-fifth of global LNG trade, moved through Hormuz in 2025; it also notes that more than 30% of global urea trade and about 20% of ammonia trade move through the Strait.",
    "S&P Global's LNG work points to a severe near-term dislocation, but its base case is a shipping suspension measured in weeks rather than years. That matters because the best ideas are the ones that still work if flows partially normalize.",
    "I therefore prefer businesses with low-cost North American nitrogen exposure or visible cost pressure on the wrong side of the energy curve, rather than names that require a 3-5 year Gulf outage to make the numbers work.",
]


IDEAS = [
    {
        "title": "Bond Long: Nutrien Ltd. 2.95% Senior Unsecured Notes due May 13, 2030",
        "ticker": "NTR",
        "security_details": [
            "CUSIP: 67077MAW8",
            "ISIN: US67077MAW82",
            "Original issue: $500 million aggregate principal amount, priced in May 2020",
            "Positioning: defensive investment-grade credit with upside from tighter nitrogen balances and pull-to-par",
        ],
        "summary": (
            "This is the cleanest bond expression of the theme. Nutrien does not need a permanent Gulf outage to make the bond work. "
            "The company already entered 2026 with stronger earnings, stronger free cash flow, and lower leverage, while its own disclosures point "
            "to improved nitrogen margins, higher upstream fertilizer volumes, and a simplified portfolio. The bond benefits from three stacked "
            "sources of return: normal carry, pull-to-par, and modest spread tightening if 2026 nitrogen conditions remain constructive."
        ),
        "why_now": [
            "Nutrien reported 2025 sales of $26.9 billion, adjusted EBITDA of $6.05 billion, and net earnings of $2.30 billion, all up versus 2024.",
            "The company generated about $4.0 billion of cash from operations in 2025 while cutting capital expenditures to roughly $2.0 billion, supporting free cash flow and balance-sheet repair.",
            "Management highlighted stronger first-quarter 2026 global urea markets, tight ammonia conditions, and improved margins at core North American nitrogen assets.",
            "The bond does not rely on an upgrade to work; it only needs Nutrien to remain a solid money-good issuer and continue de-levering.",
        ],
        "historical_headers": [
            "Year",
            "Sales ($bn)",
            "Adj. EBITDA ($bn)",
            "Net Earnings ($bn)",
            "CFO ($bn)",
        ],
        "historical_rows": [
            ["2021", "27.7", "7.1", "3.2", "3.9"],
            ["2022", "37.9", "12.2", "7.7", "8.1"],
            ["2023", "29.1", "6.1", "1.3", "5.1"],
            ["2024", "26.0", "5.4", "0.7", "3.5"],
            ["2025", "26.9", "6.0", "2.3", "4.0"],
        ],
        "projected_headers": [
            "Year",
            "Sales ($bn)",
            "Adj. EBITDA ($bn)",
            "Net Earnings ($bn)",
            "CFO ($bn)",
            "Key Assumption",
        ],
        "projected_rows": [
            ["2026E", "27.7", "6.4", "2.6", "4.3", "Constructive nitrogen and potash backdrop; no need for a prolonged Hormuz closure"],
            ["2027E", "28.4", "6.6", "2.8", "4.5", "Retail execution and stable crop input demand offset partial normalization in nitrogen"],
            ["2028E", "29.0", "6.8", "2.9", "4.7", "Normalizing market, but better mix and lower leverage preserve solid credit metrics"],
        ],
        "projection_note": (
            "Author base case. The projections assume 2026 remains constructive for global nitrogen and potash balances, but not a repeat of 2022 pricing."
        ),
        "thesis_points": [
            "Nutrien is not a pure war trade. The company already repaired earnings in 2025 and entered 2026 with better free cash flow and balance-sheet flexibility.",
            "Hormuz matters because the IEA identifies meaningful exposure of both urea and ammonia trade to the Strait. That supports nitrogen pricing even if the LNG disruption proves shorter than the market feared.",
            "Management has deliberately shifted the nitrogen mix toward core North American assets and away from weaker contributors such as Trinidad and New Madrid. That raises the quality of nitrogen cash flow.",
            "The investment case is strongest as a discounted, intermediate-duration investment-grade bond where the base case is carry and pull-to-par and the upside case is moderately tighter spreads.",
        ],
        "risks": [
            "A faster-than-expected normalization in global nitrogen pricing would reduce upside and likely leave the bond as a plain carry trade.",
            "Crop input demand can weaken if farmer economics deteriorate materially or if crop prices fall further than expected.",
            "North American natural gas costs still matter; the thesis is low-cost nitrogen, not zero-cost nitrogen.",
            "Interest-rate duration remains relevant because part of total return comes from price accretion and spread behavior rather than coupon alone.",
        ],
        "sources": [
            "https://www.nutrien.com/news/press-releases/nutrien-prices-offering-of-an-aggregate-of-ususd1-5-billion-of-3-year-10-year-and-30-year-senior-notes-1567",
            "https://www.nutrien.com/news/press-releases/nutrien-reports-fourth-quarter-and-full-year-2022-results-1669",
            "https://www.nutrien.com/news/press-releases/nutrien-reports-fourth-quarter-and-full-year-2023-results-1695",
            "https://www.nutrien.com/news/press-releases/nutrien-reports-fourth-quarter-and-full-year-2024-results-1717",
            "https://www.nutrien.com/news/press-releases/nutrien-reports-full-year-2025-results-and-provides-2026-guidance-1741",
        ],
    },
    {
        "title": "Equity Long: CF Industries Holdings, Inc.",
        "ticker": "CF",
        "security_details": [
            "Thesis type: long equity",
            "Business: low-cost North American nitrogen producer and clean-ammonia optionality",
            "What makes it fresh: cleaner pure-play nitrogen exposure than Nutrien equity, without relying on a multi-year LNG shortage",
        ],
        "summary": (
            "CF is my preferred fresh-screen long equity. It is a simpler and purer nitrogen expression than most of the original memo names, "
            "and the company has repeatedly highlighted that global nitrogen prices remain set by a cost curve anchored by higher-cost regions. "
            "That leaves a durable margin advantage for North American producers even if the Hormuz shock fades from crisis to risk premium."
        ),
        "why_now": [
            "CF's 2025 results improved to $7.08 billion of sales, $2.89 billion of adjusted EBITDA, and $1.46 billion of net earnings.",
            "Management explicitly tied 2025 pricing strength to strong global nitrogen demand and supply disruptions stemming from geopolitical issues and natural-gas availability.",
            "The company ended 2025 with $1.98 billion of cash and continues to return capital aggressively while still funding clean-energy projects.",
            "The clean-energy platform is useful but not required. The core valuation remains driven by fertilizer cash generation from a cost-advantaged North American network.",
        ],
        "historical_headers": [
            "Year",
            "Sales ($bn)",
            "Adj. EBITDA ($bn)",
            "Net Earnings ($bn)",
            "FCF ($bn)",
        ],
        "historical_rows": [
            ["2021", "6.5", "2.7", "0.9", "2.2"],
            ["2022", "11.2", "5.9", "3.4", "2.8"],
            ["2023", "6.6", "2.8", "1.5", "1.8"],
            ["2024", "5.9", "2.3", "1.2", "1.5"],
            ["2025", "7.1", "2.9", "1.5", "1.8"],
        ],
        "projected_headers": [
            "Year",
            "Sales ($bn)",
            "Adj. EBITDA ($bn)",
            "Net Earnings ($bn)",
            "FCF ($bn)",
            "Key Assumption",
        ],
        "projected_rows": [
            ["2026E", "7.4", "3.1", "1.6", "1.9", "Constructive nitrogen pricing and stable North American gas advantage"],
            ["2027E", "7.6", "3.2", "1.6", "2.0", "No 2022 spike, but structurally supportive global nitrogen cost curve"],
            ["2028E", "7.8", "3.3", "1.7", "2.0", "Steady fertilizer demand with incremental clean-energy optionality"],
        ],
        "projection_note": (
            "Author base case. The estimates assume supportive but not peak-cycle nitrogen markets and continued disciplined buybacks."
        ),
        "thesis_points": [
            "CF has the cleanest direct exposure to nitrogen margins among large-cap U.S. names in this theme.",
            "Its own 2025 release explicitly says lower global operating rates and natural-gas constraints in other regions supported the global market clearing price.",
            "The balance sheet and cash generation are good enough that the company does not need heroic fertilizer pricing to keep buying back stock and funding growth.",
            "If the market increasingly values North American low-cost ammonia and nitrogen as strategic assets, CF should outperform more diversified peers.",
        ],
        "risks": [
            "Nitrogen prices are volatile and can normalize quickly if export availability improves or import demand slows.",
            "Higher North American natural-gas costs would compress the cost advantage if not offset by product pricing.",
            "Clean-ammonia projects involve execution, permitting, and customer-commitment risk.",
            "Because CF is a purer nitrogen name, it can underperform diversified fertilizer peers if nitrogen-specific fundamentals turn down.",
        ],
        "sources": [
            "https://www.cfindustries.com/newsroom/2022/fullyear2021results",
            "https://www.cfindustries.com/newsroom/2023/fullyear2022earnings",
            "https://www.cfindustries.com/newsroom/2024/q4-2023-earnings",
            "https://www.cfindustries.com/newsroom/2025/full-year-2024-results",
            "https://www.cfindustries.com/newsroom/2026/fy-2025-earnings",
        ],
    },
    {
        "title": "Equity Short: O-I Glass, Inc.",
        "ticker": "OI",
        "security_details": [
            "Thesis type: short equity",
            "Business: energy-intensive glass packaging manufacturer with significant European exposure",
            "What makes it fresh: cleaner expression of higher European energy costs than airlines, freight forwarders, or tourism names",
        ],
        "summary": (
            "O-I is my preferred short because the pressure points are explicit and operational. Management guided to higher 2026 EBITDA despite "
            "an estimated $150 million energy-cost step-up after favorable European energy contracts expired at year-end 2025. I think that guidance "
            "still leaves too much credit for volume recovery and cost absorption in a muted demand environment."
        ),
        "why_now": [
            "O-I's 2025 adjusted earnings improved to $1.60 per share, but net debt remained about $4.2 billion and the company remains highly sensitive to execution and energy costs.",
            "Management's 2026 guide calls for adjusted EBITDA of $1.25 billion to $1.30 billion and approximately $200 million of free cash flow even after the European energy reset.",
            "That guide embeds meaningful Fit To Win benefits and operational discipline. If volume stays soft, Europe remains competitive on price, or energy remains elevated, the miss path is straightforward.",
            "Unlike many consumer-travel shorts, this is not mainly a traffic timing call. It is a cost-curve and fixed-asset utilization call.",
        ],
        "historical_headers": [
            "Year",
            "Sales ($bn)",
            "Seg. Op. Profit ($bn)",
            "Adj. EPS ($)",
            "FCF ($bn)",
            "Net Debt ($bn)",
        ],
        "historical_rows": [
            ["2021", "6.4", "0.83", "1.83", "0.28", "4.1"],
            ["2022", "6.9", "0.96", "2.30", "0.24", "3.9"],
            ["2023", "7.1", "1.19", "3.09", "0.13", "4.0"],
            ["2024", "6.5", "0.75", "0.81", "-0.13", "4.2"],
            ["2025", "6.4", "0.85", "1.60", "0.17", "4.2"],
        ],
        "projected_headers": [
            "Year",
            "Sales ($bn)",
            "Seg. Op. Profit ($bn)",
            "Adj. EPS ($)",
            "FCF ($bn)",
            "Net Debt ($bn)",
            "Key Assumption",
        ],
        "projected_rows": [
            ["2026E", "6.2", "0.78", "1.35", "0.13", "4.3", "Energy reset and weak demand keep results below management's EBITDA guide"],
            ["2027E", "6.2", "0.76", "1.25", "0.10", "4.3", "Some savings, but no strong recovery in Europe or beverage demand"],
            ["2028E", "6.1", "0.74", "1.20", "0.10", "4.3", "Mature business with persistent pricing pressure and limited deleveraging"],
        ],
        "projection_note": (
            "Author bear/base case for the short. Management's 2026 guidance is higher than my estimate and is therefore the main valuation challenge."
        ),
        "thesis_points": [
            "O-I's 2026 setup includes a visible and quantified energy headwind rather than a speculative one.",
            "Glass production remains capital intensive and fixed-cost heavy. When demand is soft, the combination of idle capacity, weaker pricing, and higher energy costs hurts quickly.",
            "The balance sheet has not repaired enough to make misses painless. Net debt remained about $4.2 billion at the end of 2025 and leverage was still 3.5x.",
            "As a short, O-I is better than tourism or logistics shorts because the company itself has already laid out the cost reset and the execution burden.",
        ],
        "risks": [
            "Fit To Win savings could exceed expectations and allow management to offset more of the energy step-up than I expect.",
            "European energy prices could fall faster than assumed, particularly if LNG flows normalize quickly and inventories rebuild.",
            "A stronger beverage and food packaging recovery could lift utilization and absorb fixed costs better than expected.",
            "Short interest can become crowded if the market treats O-I as a restructuring self-help story rather than a volume-and-energy story.",
        ],
        "sources": [
            "https://www.o-i.com/news/oi-glass-reports-full-year-and-fourth-quarter-2021-results/",
            "https://www.o-i.com/news/o-i-glass-reports-full-year-and-fourth-quarter-2022-results/",
            "https://www.o-i.com/news/o-i-glass-reports-full-year-and-fourth-quarter-2023-results/",
            "https://www.o-i.com/news/o-i-glass-reports-full-year-and-fourth-quarter-2024-results/",
            "https://www.o-i.com/news/o-i-glass-reports-full-year-and-fourth-quarter-2025-results/",
        ],
    },
]


PAIR_TRADE = {
    "title": "Preferred Relative-Value Expression: Long CF Industries / Short O-I Glass",
    "points": [
        "Long the low-cost North American nitrogen producer whose economics improve when the global nitrogen cost curve tightens.",
        "Short the energy-intensive glass producer facing a visible European cost reset and a less forgiving fixed-cost model.",
        "This pair does not require a permanent Hormuz closure. It only needs the shock to leave behind tighter fertilizer balances and still-high enough European energy costs to pressure O-I's guidance.",
        "The pair also reduces outright market beta compared with running a standalone cyclical long or a standalone industrial short.",
    ],
    "disconfirmants": [
        "Rapid normalization in nitrogen prices with no lasting benefit to North American fertilizer producers.",
        "A larger-than-expected fall in European power and gas costs combined with stronger beverage packaging demand.",
        "Evidence that O-I's self-help initiatives are structurally outweighing the energy reset.",
    ],
}


SOURCE_APPENDIX = [
    "Macro / market structure",
    "https://www.iea.org/topics/the-middle-east-and-global-energy-markets",
    "https://www.spglobal.com/market-intelligence/en/news-insights/research/2026/03/middle-east-war-disrupts-global-lng-near-term-outlook",
    "",
    "Equity prices cited in working notes came from finance data as of the March 25, 2026 U.S. close. Bond pricing should be refreshed from TRACE or a trading platform immediately before use.",
]


def set_document_defaults(document: Document) -> None:
    styles = document.styles
    styles["Normal"].font.name = "Calibri"
    styles["Normal"].font.size = Pt(10.5)
    styles["Heading 1"].font.name = "Calibri"
    styles["Heading 1"].font.size = Pt(16)
    styles["Heading 2"].font.name = "Calibri"
    styles["Heading 2"].font.size = Pt(12.5)
    styles["Heading 3"].font.name = "Calibri"
    styles["Heading 3"].font.size = Pt(11)

    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)


def add_title(document: Document) -> None:
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Fresh-Screen Hormuz Trade Ideas")
    r.bold = True
    r.font.size = Pt(20)

    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Bond long and equity long/short ideas rebuilt from first principles")
    r.italic = True
    r.font.size = Pt(11)

    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Date: March 26, 2026\n")
    p.add_run("Prepared for Tim").bold = True


def add_bullets(document: Document, items):
    for item in items:
        document.add_paragraph(item, style="List Bullet")


def add_table(document: Document, headers, rows, first_col_width=None):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = True
    hdr_cells = table.rows[0].cells
    for i, header in enumerate(headers):
        hdr_cells[i].text = header
        for paragraph in hdr_cells[i].paragraphs:
            for run in paragraph.runs:
                run.bold = True
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
    if first_col_width:
        for row in table.rows:
            row.cells[0].width = first_col_width
    document.add_paragraph()


def add_sources(document: Document, sources):
    for source in sources:
        document.add_paragraph(source, style="List Bullet")


def add_idea_section(document: Document, idea: dict) -> None:
    document.add_heading(idea["title"], level=1)
    document.add_paragraph(idea["summary"])

    document.add_heading("Security / setup", level=2)
    add_bullets(document, idea["security_details"])

    document.add_heading("Why this idea survives a shorter disruption", level=2)
    add_bullets(document, idea["why_now"])

    document.add_heading("Five-year historical financials", level=2)
    add_table(document, idea["historical_headers"], idea["historical_rows"], first_col_width=Inches(0.75))

    document.add_heading("Projected financials", level=2)
    add_table(document, idea["projected_headers"], idea["projected_rows"], first_col_width=Inches(0.75))
    note = document.add_paragraph()
    note_run = note.add_run("Projection note: ")
    note_run.bold = True
    note.add_run(idea["projection_note"])

    document.add_heading("Detailed thesis", level=2)
    add_bullets(document, idea["thesis_points"])

    document.add_heading("Key risks", level=2)
    add_bullets(document, idea["risks"])

    document.add_heading("Primary sources", level=2)
    add_sources(document, idea["sources"])


def build_document() -> Document:
    document = Document()
    set_document_defaults(document)
    add_title(document)

    document.add_paragraph()
    document.add_heading("Executive summary", level=1)
    document.add_paragraph(
        "This memo rebuilds the trade ideas from scratch. The core conclusion is that the best expressions of the "
        "Hormuz shock are the ones that still work if the disruption proves severe but temporary. I therefore prefer "
        "a discounted investment-grade bond issued by Nutrien, a clean nitrogen long in CF Industries, and a short in "
        "O-I Glass, which faces a visible European energy reset and fixed-cost pressure."
    )

    document.add_heading("Macro frame", level=1)
    add_bullets(document, MACRO_POINTS)

    document.add_heading("What I would actually own", level=1)
    add_bullets(
        document,
        [
            "Bond long: Nutrien Ltd. 2.95% due May 13, 2030.",
            "Equity long: CF Industries.",
            "Equity short: O-I Glass.",
            "Preferred pair trade: Long CF / Short O-I.",
        ],
    )

    for idea in IDEAS:
        document.add_section(WD_SECTION.NEW_PAGE)
        add_idea_section(document, idea)

    document.add_section(WD_SECTION.NEW_PAGE)
    document.add_heading(PAIR_TRADE["title"], level=1)
    add_bullets(document, PAIR_TRADE["points"])

    document.add_heading("Main disconfirmants", level=2)
    add_bullets(document, PAIR_TRADE["disconfirmants"])

    document.add_heading("Source appendix", level=1)
    for item in SOURCE_APPENDIX:
        if item:
            document.add_paragraph(item, style="List Bullet")
        else:
            document.add_paragraph()

    document.add_paragraph()
    disclaimer = document.add_paragraph()
    disclaimer_run = disclaimer.add_run(
        "Note: projected financials are author estimates as of March 26, 2026 and are not company guidance unless explicitly stated."
    )
    disclaimer_run.italic = True

    return document


def main() -> None:
    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    document = build_document()
    document.save(OUTPUT_PATH)
    print(f"Wrote {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
